"""The live-formula Excel export.

Two kinds of check here.

The fast ones assert structure: that the workbook ships FORMULAS rather
than pasted values, and that each Excel chart default known to render a
broken chart has been explicitly overridden. Those defaults are invisible
to XML inspection -- an element you never wrote has no XML to inspect --
so asserting that the attribute is set is the only mechanical check
available.

The slow one evaluates the workbook through LibreOffice and compares every
headline figure against the engine. It is skipped where LibreOffice is not
installed.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

import evm
from evm.excel_export import export_workbook

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def exported(project, tmp_path_factory) -> Path:
    destination = tmp_path_factory.mktemp("export") / "EVM_Report.xlsx"
    return export_workbook(project, destination)


@pytest.fixture(scope="module")
def book(exported):
    return openpyxl.load_workbook(exported)


@pytest.fixture(scope="module")
def evaluated(exported):
    """The exported workbook with every formula computed by LibreOffice."""
    sys.path.insert(0, str(ROOT / "scripts"))
    from validate_benchmark import recalculate
    return openpyxl.load_workbook(recalculate(exported), data_only=True)


class TestStructure:
    def test_expected_sheets_are_present(self, book):
        assert set(book.sheetnames) == {
            "Dashboard", "Inputs", "Baseline", "EVM", "Variance", "Read Me"
        }

    def test_dashboard_is_the_first_sheet(self, book):
        """It is the sheet the reader should land on."""
        assert book.sheetnames[0] == "Dashboard"

    def test_workbook_recalculates_on_open(self, book):
        """openpyxl writes no cached results, so without this flag a reader
        that trusts the cache sees blank cells."""
        assert book.calculation.fullCalcOnLoad is True


class TestItShipsFormulasNotValues:
    """The headline claim of the export: it is live, not a snapshot."""

    @pytest.mark.parametrize("cell", ["C5", "C6", "C7", "C8"])
    def test_dashboard_position_cells_are_formulas(self, book, cell):
        value = book["Dashboard"][cell].value
        assert isinstance(value, str) and value.startswith("=")

    def test_every_evm_index_is_a_formula(self, book):
        ws = book["EVM"]
        for row in range(5, 17):
            for col in "CDEFGHIJK":
                value = ws[f"{col}{row}"].value
                assert isinstance(value, str) and value.startswith("="), f"{col}{row}"

    def test_earned_value_is_computed_from_budget_times_percent(self, book):
        """EV must be derived in the sheet, not carried across from Python."""
        formula = book["EVM"]["D5"].value
        assert "SUMPRODUCT" in formula
        assert "Inputs!" in formula

    def test_the_verdict_is_a_formula_so_it_tracks_edits(self, book):
        """A pasted verdict would still read ON TRACK after the user edits
        the inputs to something catastrophic."""
        ws = book["Dashboard"]
        found = [
            ws.cell(row=r, column=3).value
            for r in range(1, 40)
            if str(ws.cell(row=r, column=2).value) == "VERDICT"
        ]
        assert found and isinstance(found[0], str) and found[0].startswith("=IF(")

    def test_inputs_are_values_since_they_are_the_editable_data(self, book):
        ws = book["Inputs"]
        for row in range(5, 24):
            assert isinstance(ws[f"C{row}"].value, (int, float))


class TestChartDefaultsAreOverridden:
    """Each of these defaults produces a visibly broken chart.

    None of them can be caught by reading back the XML we wrote, because
    the failure is an element we did NOT write being defaulted by Excel.
    Asserting the override is set is the only mechanical check there is;
    appearance still needs a human to open the file.
    """

    def test_category_axes_are_pinned_to_text(self, book):
        """Date-like labels promote a category axis to a DATE axis, which
        lays out a continuous calendar and collapses bars to hairlines."""
        for chart in book["Dashboard"]._charts + book["Variance"]._charts:
            assert chart.x_axis.auto is False

    def test_category_labels_are_pinned_to_the_base_of_the_plot(self, book):
        """tickLblPos defaults to the axis CROSSING, which on all-negative
        data is the top of the plot, over the chart title."""
        for chart in book["Dashboard"]._charts + book["Variance"]._charts:
            assert chart.x_axis.tickLblPos == "low"

    def test_negative_bars_are_not_inverted_to_white(self, book):
        """An absent invertIfNegative is treated as TRUE, so an
        all-negative cost variance series renders as hollow outlines."""
        for chart in book["Variance"]._charts:
            for series in chart.series:
                assert series.invertIfNegative is False

    def test_series_titles_are_set_explicitly(self, book):
        """titles_from_data pulls the header cell verbatim, including any
        annotation, straight into the legend."""
        for chart in book["Dashboard"]._charts:
            for series in chart.series:
                assert series.tx is not None
                assert series.tx.v

    def test_index_axis_is_pinned_so_periods_are_comparable(self, book):
        """An auto-scaled index axis makes the same CPI look catastrophic
        or negligible depending on the rest of the data."""
        trend = [c for c in book["Dashboard"]._charts
                 if c.title is not None and "Performance Trend" in str(c.title)]
        assert trend, "performance trend chart not found"
        assert trend[0].y_axis.scaling.min == 0.5
        assert trend[0].y_axis.scaling.max == 1.2


class TestSeriesStopAtTheDataDate:
    def test_earned_value_returns_na_beyond_the_data_date(self, book, project):
        """A formula returning "" is plotted as ZERO, not as a gap, so the
        earned value line would dive to the axis after the data date.
        NA() is skipped by the chart."""
        ws = book["EVM"]
        for period in range(project.data_date + 1, project.n_periods + 1):
            row = 5 + period - 1
            assert ws[f"D{row}"].value == "=NA()"
            assert ws[f"E{row}"].value == "=NA()"

    def test_planned_value_runs_the_full_baseline(self, book, project):
        """PV is a plan; it is known for every period and must be drawn."""
        ws = book["EVM"]
        for period in range(1, project.n_periods + 1):
            row = 5 + period - 1
            assert ws[f"C{row}"].value.startswith("=SUM(Baseline!")


@pytest.mark.skipif(
    shutil.which("soffice") is None
    and not Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists(),
    reason="LibreOffice is needed to evaluate the exported workbook",
)
class TestExportedFiguresMatchTheEngine:
    """Evaluate the export and compare it with the engine.

    This closes the loop: the dashboard, the manual benchmark and the
    exported workbook all agree.
    """

    def _dashboard(self, book):
        ws = book["Dashboard"]
        return {
            str(ws.cell(row=r, column=2).value): ws.cell(row=r, column=3).value
            for r in range(1, 40)
            if ws.cell(row=r, column=2).value
        }

    def test_base_measures_match(self, evaluated, project):
        values = self._dashboard(evaluated)
        m = evm.metrics_at(project, project.data_date)
        assert values["Planned value (PV)"] == pytest.approx(m.pv, abs=0.01)
        assert values["Earned value (EV)"] == pytest.approx(m.ev, abs=0.01)
        assert values["Actual cost (AC)"] == pytest.approx(m.ac, abs=0.01)

    def test_indices_match(self, evaluated, project):
        values = self._dashboard(evaluated)
        m = evm.metrics_at(project, project.data_date)
        assert values["Cost performance index (CPI)"] == pytest.approx(m.cpi, abs=1e-9)
        assert values["Schedule performance index (SPI)"] == pytest.approx(m.spi, abs=1e-9)

    def test_forecasts_match(self, evaluated, project):
        values = self._dashboard(evaluated)
        m = evm.metrics_at(project, project.data_date)
        eac = {f.method: f for f in evm.forecasts(m)}
        assert values["EAC — performance continues"] == pytest.approx(
            eac["EAC1"].eac, abs=0.01
        )
        assert values["EAC — variance was a one-off"] == pytest.approx(
            eac["EAC2"].eac, abs=0.01
        )
        assert values["EAC — cost and schedule pressure"] == pytest.approx(
            eac["EAC3"].eac, abs=0.01
        )
        assert values["TCPI to finish on budget"] == pytest.approx(
            evm.tcpi_to_bac(m), abs=1e-9
        )

    def test_verdict_reflects_the_case_study(self, evaluated):
        values = self._dashboard(evaluated)
        assert values["VERDICT"] == "BEHIND SCHEDULE AND OVERSPENT"
