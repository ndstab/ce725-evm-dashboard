"""Compare the Python engine against the manual Excel benchmark.

The comparison is between two independently written implementations of the
same definitions:

    the evm package            Python
    EVM_Manual_Benchmark.xlsx  typed Excel formulas

Critically, the Excel side is EVALUATED, not re-implemented. The workbook
is driven through LibreOffice in headless mode, which loads it, recalculates
every formula with a real spreadsheet engine and writes the results back.
We then read those results.

This distinction is the whole point. A validation script that re-implements
the intended logic in Python and compares it with the engine is checking
our understanding against itself: if we misunderstood the formula, both
sides misunderstand it identically and the test passes. Executing the
workbook removes that shared assumption. (The idea is taken from the
verification approach in Tahir-yamin/excel-scurve-builder, credited in the
method note.)

Outputs
    output/benchmark-recalculated.xlsx   workbook with computed values
    docs/generated/benchmark-values.json snapshot the fast test reads
    docs/generated/benchmark-comparison.md  the discrepancy table

Run:  python scripts/validate_benchmark.py
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import evm  # noqa: E402

BENCHMARK = ROOT / "benchmark" / "EVM_Manual_Benchmark.xlsx"
OUTPUT_DIR = ROOT / "output"
GENERATED_DIR = ROOT / "docs" / "generated"

BENCHMARK_PERIODS = (5, 6)

# Tolerances. Rupee figures are compared to the paisa; dimensionless indices
# to nine decimal places. Both are far tighter than anything that could
# matter to a decision -- the point is to detect a formula difference, not
# to assess materiality.
RUPEE_TOL = 0.01
RATIO_TOL = 1e-9


class LibreOfficeUnavailable(RuntimeError):
    pass


def _soffice() -> str:
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if mac.exists():
        return str(mac)
    raise LibreOfficeUnavailable(
        "LibreOffice was not found. It is needed to evaluate the benchmark "
        "workbook's formulas. Install it from libreoffice.org, or open the "
        "workbook in Excel and save it, which also caches the computed values."
    )


def recalculate(workbook: Path) -> Path:
    """Evaluate a workbook's formulas and return a copy holding the results.

    The workbook is written by openpyxl, which stores formulas and no cached
    results, and is flagged fullCalcOnLoad. LibreOffice therefore computes
    every cell on open and caches it on save.
    """
    soffice = _soffice()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmp:
        # A dedicated profile keeps this from colliding with a LibreOffice
        # window the user already has open, which otherwise makes the
        # headless call exit silently without converting anything.
        profile = Path(tmp) / "profile"
        result = subprocess.run(
            [
                soffice, "--headless", "--norestore", "--invisible",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to", "xlsx", "--outdir", tmp, str(workbook),
            ],
            capture_output=True, text=True, timeout=300,
        )
        produced = Path(tmp) / workbook.name
        if not produced.exists():
            raise RuntimeError(
                "LibreOffice did not produce a recalculated workbook.\n"
                f"stdout: {result.stdout}\nstderr: {result.stderr}"
            )
        destination = OUTPUT_DIR / "benchmark-recalculated.xlsx"
        shutil.copy(produced, destination)
    return destination


def read_benchmark(recalculated: Path) -> dict[str, dict[int, float]]:
    """Read the computed Indices sheet into {measure: {period: value}}."""
    wb = openpyxl.load_workbook(recalculated, data_only=True)
    ws = wb["Indices"]
    values: dict[str, dict[int, float]] = {}
    for row in range(5, 200):
        label = ws.cell(row=row, column=2).value
        if not label:
            break
        entry = {}
        for col, period in zip((4, 5), BENCHMARK_PERIODS):
            cell = ws.cell(row=row, column=col).value
            if cell is None or isinstance(cell, str):
                raise RuntimeError(
                    f"'{label}' period {period} did not evaluate to a number "
                    f"(got {cell!r}). The workbook was not recalculated."
                )
            entry[period] = float(cell)
        values[str(label)] = entry
    if not values:
        raise RuntimeError("the Indices sheet was empty")
    return values


def engine_values(project: "evm.Project", period: int) -> dict[str, tuple[float, str]]:
    """Every measure the benchmark reports, computed by the Python engine.

    Returned as {label: (value, kind)} where kind selects the tolerance.
    Labels match the benchmark workbook exactly so the two can be zipped.
    """
    m = evm.metrics_at(project, period)
    es = evm.earned_schedule(project, period)
    eac = {f.method: f for f in evm.forecasts(m)}

    return {
        "Budget at completion (BAC)": (m.bac, "rupee"),
        "Planned value (PV)": (m.pv, "rupee"),
        "Earned value (EV)": (m.ev, "rupee"),
        "Actual cost (AC)": (m.ac, "rupee"),
        "Cost variance (CV)": (m.cv, "rupee"),
        "Schedule variance (SV)": (m.sv, "rupee"),
        # The workbook holds these as Excel percentages, i.e. fractions.
        "Cost variance %": (m.cv_pct / 100.0, "ratio"),
        "Schedule variance %": (m.sv_pct / 100.0, "ratio"),
        "Cost performance index (CPI)": (m.cpi, "ratio"),
        "Schedule performance index (SPI)": (m.spi, "ratio"),
        "Percent complete": (m.pct_complete / 100.0, "ratio"),
        "Percent spent": (m.pct_spent / 100.0, "ratio"),
        "Percent scheduled": (m.pct_scheduled / 100.0, "ratio"),
        "Earned schedule (ES)": (es.earned_schedule, "ratio"),
        "Actual time (AT)": (es.actual_time, "ratio"),
        "Schedule variance, time SV(t)": (es.sv_t, "ratio"),
        "Schedule performance index SPI(t)": (es.spi_t, "ratio"),
        "EAC method 1": (eac["EAC1"].eac, "rupee"),
        "EAC method 2": (eac["EAC2"].eac, "rupee"),
        "EAC method 3": (eac["EAC3"].eac, "rupee"),
        "EAC method 4": (eac["EAC4"].eac, "rupee"),
        "ETC (on EAC method 1)": (eac["EAC1"].etc, "rupee"),
        "VAC (on EAC method 1)": (eac["EAC1"].vac, "rupee"),
        "TCPI to BAC": (evm.tcpi_to_bac(m), "ratio"),
        "TCPI to EAC1": (eac["EAC1"].tcpi, "ratio"),
    }


def compare(project, benchmark: dict[str, dict[int, float]]) -> list[dict]:
    """Row per measure per period, with the difference and a verdict."""
    rows = []
    for period in BENCHMARK_PERIODS:
        for label, (tool_value, kind) in engine_values(project, period).items():
            if label not in benchmark:
                rows.append({
                    "period": period, "measure": label, "tool": tool_value,
                    "benchmark": None, "difference": None, "kind": kind,
                    "match": False, "note": "not present in the workbook",
                })
                continue
            bench_value = benchmark[label][period]
            difference = tool_value - bench_value
            tolerance = RUPEE_TOL if kind == "rupee" else RATIO_TOL
            rows.append({
                "period": period, "measure": label, "tool": tool_value,
                "benchmark": bench_value, "difference": difference,
                "kind": kind, "match": abs(difference) <= tolerance,
                "note": "" if abs(difference) <= tolerance else "EXCEEDS TOLERANCE",
            })
    return rows


def _fmt(value: float | None, kind: str) -> str:
    if value is None:
        return "—"
    if kind == "rupee":
        return f"{value:,.2f}"
    return f"{value:.9f}"


def write_report(rows: list[dict], destination: Path) -> None:
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    matched = sum(1 for r in rows if r["match"])

    lines = [
        "<!-- GENERATED FILE. Produced by scripts/validate_benchmark.py. Do not edit. -->",
        "",
        "# Tool versus manual benchmark — full comparison",
        "",
        "The tool column is the Python engine. The benchmark column is the",
        "typed Excel formula in `benchmark/EVM_Manual_Benchmark.xlsx`, evaluated",
        "by LibreOffice rather than re-implemented in Python.",
        "",
        f"Tolerance: {RUPEE_TOL} rupee on currency, {RATIO_TOL:g} on dimensionless indices.",
        "",
        f"**{matched} of {len(rows)} comparisons agree within tolerance.**",
        "",
    ]
    for period in BENCHMARK_PERIODS:
        lines += [
            f"## Reporting period {period}",
            "",
            "| Measure | Tool | Manual benchmark | Difference | Agrees |",
            "|---|---:|---:|---:|:---:|",
        ]
        for row in [r for r in rows if r["period"] == period]:
            kind = row["kind"]
            lines.append(
                f"| {row['measure']} | {_fmt(row['tool'], kind)} | "
                f"{_fmt(row['benchmark'], kind)} | "
                f"{_fmt(row['difference'], kind)} | "
                f"{'yes' if row['match'] else '**NO**'} |"
            )
        lines.append("")
    destination.write_text("\n".join(lines))


def main() -> int:
    project = evm.load_project(ROOT / "data")

    if not BENCHMARK.exists():
        print(f"Benchmark workbook not found at {BENCHMARK}.")
        print("Build it first:  python benchmark/build_benchmark.py")
        return 2

    print("Evaluating the benchmark workbook through LibreOffice...")
    recalculated = recalculate(BENCHMARK)
    benchmark = read_benchmark(recalculated)
    print(f"  read {len(benchmark)} measures for periods {BENCHMARK_PERIODS}")

    rows = compare(project, benchmark)

    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    snapshot = GENERATED_DIR / "benchmark-values.json"
    snapshot.write_text(json.dumps(
        {label: {str(p): v for p, v in periods.items()}
         for label, periods in benchmark.items()},
        indent=2, sort_keys=True,
    ) + "\n")

    report = GENERATED_DIR / "benchmark-comparison.md"
    write_report(rows, report)

    failures = [r for r in rows if not r["match"]]
    width = max(len(r["measure"]) for r in rows)
    for period in BENCHMARK_PERIODS:
        print(f"\nPeriod {period}")
        for row in [r for r in rows if r["period"] == period]:
            flag = "ok " if row["match"] else "DIFF"
            print(
                f"  {flag}  {row['measure']:<{width}}  "
                f"tool {_fmt(row['tool'], row['kind']):>22}  "
                f"benchmark {_fmt(row['benchmark'], row['kind']):>22}"
            )

    print(f"\n{len(rows) - len(failures)} of {len(rows)} comparisons agree.")
    print(f"Snapshot : {snapshot.relative_to(ROOT)}")
    print(f"Report   : {report.relative_to(ROOT)}")
    if failures:
        print(f"\n{len(failures)} DISCREPANCIES:")
        for row in failures:
            print(f"  period {row['period']}  {row['measure']}  "
                  f"difference {row['difference']}  {row['note']}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
